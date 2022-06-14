from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
# #B2:C11 데이터로 차트생성
# bar_value = Reference(ws, min_row=2, max_row=11,min_col=2, max_col=3)
# bar_chart=BarChart() # chart 종류( 바, 라인,파이...)
# bar_chart.add_data(bar_value) # 차트 데이터 추가

# ws.add_chart(bar_chart, "E1") # 차트 위치 지장

# B1:C11 제목 포함 데이터
line_value= Reference(ws, min_row=1, max_row=11,min_col=2, max_col=3)
line_chart= LineChart()
line_chart.add_data(line_value, titles_from_data= True) #계열 영어 수학 from 데이타
line_chart.title= "성적표"
line_chart.style= 10 # 미리 정해진 스타일,사용자정의도 가능
line_chart.y_axis.title="점수"
line_chart.x_axis.title="번호"
ws.add_chart(line_chart, "E1")


wb.save("sample_chart.xlsx")